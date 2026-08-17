import os
import tempfile

from aiogram import Router, F
from aiogram.types import Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton
from aiogram.types.input_file import FSInputFile
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload
from database import async_session
from models import User, Task, TaskApplication, ApplicationStatus, Division
from config import rank_name, get_monthly_coefficient
from handlers.admin import is_admin
from datetime import datetime

router = Router()


class PeriodReport(StatesGroup):
    date_from = State()
    date_to = State()


def _parse_date(s: str) -> datetime | None:
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s.strip(), fmt)
        except ValueError:
            continue
    return None


async def _calc_worker_pay_monthly(worker: User, month: int | None = None, year: int | None = None) -> tuple[float, int, float, float]:
    now = datetime.utcnow()
    if month is None:
        month = now.month
    if year is None:
        year = now.year

    month_start = f"{year:04d}-{month:02d}-01"
    if month == 12:
        month_end = f"{year+1:04d}-01-01"
    else:
        month_end = f"{year:04d}-{month+1:02d}-01"

    async with async_session() as session:
        confirmed_apps = (await session.execute(
            select(TaskApplication)
            .options(selectinload(TaskApplication.task))
            .where(
                TaskApplication.user_id == worker.id,
                TaskApplication.status == ApplicationStatus.CONFIRMED.value,
            )
        )).scalars().all()

        coeff_divs = (await session.execute(
            select(Division.id).where(Division.counts_for_coeff == True)
        )).scalars().all()
        coeff_div_ids = set(coeff_divs)

        penalties = (await session.execute(
            select(func.coalesce(func.sum(TaskApplication.penalty_amount), 0))
            .where(
                TaskApplication.user_id == worker.id,
                TaskApplication.status == ApplicationStatus.PENALTY.value,
            )
        )).scalar() or 0

    month_apps = []
    for app in confirmed_apps:
        task_date = _parse_task_date(app.task.date)
        if task_date and month_start <= task_date.strftime("%Y-%m-%d") < month_end:
            month_apps.append(app)

    coeff_count = sum(
        1 for app in month_apps
        if app.task.division_id and app.task.division_id in coeff_div_ids
    )

    coeff = get_monthly_coefficient(coeff_count)
    base_sum = sum(app.task.rate_for_rank(worker.rank) for app in month_apps)
    earned = base_sum * coeff - penalties

    return earned, len(month_apps), penalties, coeff


async def _calc_worker_pay_period(worker: User, date_from: str, date_to: str) -> tuple[float, int, float, float]:
    async with async_session() as session:
        confirmed_apps = (await session.execute(
            select(TaskApplication)
            .options(selectinload(TaskApplication.task))
            .where(
                TaskApplication.user_id == worker.id,
                TaskApplication.status == ApplicationStatus.CONFIRMED.value,
            )
        )).scalars().all()

        coeff_divs = (await session.execute(
            select(Division.id).where(Division.counts_for_coeff == True)
        )).scalars().all()
        coeff_div_ids = set(coeff_divs)

        penalties = (await session.execute(
            select(func.coalesce(func.sum(TaskApplication.penalty_amount), 0))
            .where(
                TaskApplication.user_id == worker.id,
                TaskApplication.status == ApplicationStatus.PENALTY.value,
            )
        )).scalar() or 0

    period_apps = []
    for app in confirmed_apps:
        task_date = _parse_task_date(app.task.date)
        if task_date and date_from <= task_date.strftime("%Y-%m-%d") <= date_to:
            period_apps.append(app)

    coeff_count = sum(
        1 for app in period_apps
        if app.task.division_id and app.task.division_id in coeff_div_ids
    )

    coeff = get_monthly_coefficient(coeff_count)
    base_sum = sum(app.task.rate_for_rank(worker.rank) for app in period_apps)
    earned = base_sum * coeff - penalties

    return earned, len(period_apps), penalties, coeff


def _parse_task_date(date_str: str) -> datetime | None:
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(date_str.strip(), fmt)
        except ValueError:
            continue
    return None


# --- Расчёт оплаты (за текущий месяц) ---

@router.message(Command("payroll"))
async def cmd_payroll(message: Message):
    if not await is_admin(message.from_user.id):
        await message.answer("У вас нет прав администратора.")
        return

    async with async_session() as session:
        workers = (await session.execute(
            select(User).where(User.role == "worker")
        )).scalars().all()

    if not workers:
        await message.answer("Работников нет.")
        return

    now = datetime.utcnow()
    month_name = now.strftime("%B %Y")
    lines = []
    total = 0.0

    for worker in workers:
        earned, confirmed_count, penalties, coeff = await _calc_worker_pay_monthly(worker)
        lines.append(
            f"• {worker.full_name} | {rank_name(worker.rank)}\n"
            f"  Задач (месяц): {confirmed_count} | Коэфф: x{coeff}\n"
            f"  Штрафы: -{penalties:.0f}₽\n"
            f"  Итого: {earned:.0f}₽"
        )
        total += earned

    report = f"💰 Расчёт оплаты ({now.strftime('%m.%Y')}):\n\n" + "\n\n".join(lines)
    report += f"\n\n{'='*30}\nОбщий итог: {total:.0f}₽"

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📊 Отчёт за период", callback_data="period_report")]
    ])
    await message.answer(report, reply_markup=kb)


# --- Экспорт Excel (за текущий месяц) ---

@router.message(Command("payroll_xlsx"))
async def cmd_payroll_xlsx(message: Message):
    if not await is_admin(message.from_user.id):
        await message.answer("У вас нет прав администратора.")
        return

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    async with async_session() as session:
        workers = (await session.execute(
            select(User).where(User.role == "worker")
        )).scalars().all()

    if not workers:
        await message.answer("Работников нет.")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Расчёт оплаты"

    headers = ["ФИО", "Ранг", "Роль", "Задач (мес.)", "Коэфф.", "Штрафы ₽", "Итого ₽"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    total = 0.0
    row = 2
    for worker in workers:
        earned, confirmed_count, penalties, coeff = await _calc_worker_pay_monthly(worker)
        roles = ", ".join(worker.job_roles) if worker.job_roles else worker.job_role
        ws.cell(row=row, column=1, value=worker.full_name)
        ws.cell(row=row, column=2, value=rank_name(worker.rank))
        ws.cell(row=row, column=3, value=roles)
        ws.cell(row=row, column=4, value=confirmed_count)
        ws.cell(row=row, column=5, value=coeff)
        ws.cell(row=row, column=6, value=round(penalties))
        ws.cell(row=row, column=7, value=round(earned))
        total += earned
        row += 1

    ws.cell(row=row + 1, column=6, value="ИТОГО:").font = Font(bold=True)
    ws.cell(row=row + 1, column=7, value=round(total)).font = Font(bold=True)

    for col in ws.columns:
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col[0].column_letter].width = max_len + 3

    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)

    date_str = datetime.utcnow().strftime("%d.%m.%Y")
    doc = FSInputFile(path, filename=f"payroll_{date_str}.xlsx")
    await message.answer_document(doc, caption=f"💰 Расчёт оплаты за {date_str}")
    os.unlink(path)


# --- Отчёт за период ---

@router.callback_query(F.data == "period_report")
async def cb_period_report(callback: CallbackQuery, state: FSMContext):
    if not await is_admin(callback.from_user.id):
        return
    await state.set_state(PeriodReport.date_from)
    await callback.message.answer("Введите дату начала периода (например, 01.08.2026):")
    await callback.answer()


@router.message(PeriodReport.date_from)
async def period_date_from(message: Message, state: FSMContext):
    d = _parse_date(message.text)
    if not d:
        await message.answer("Неверный формат. Введите дату (например, 01.08.2026):")
        return
    await state.update_data(date_from=d.strftime("%Y-%m-%d"), date_from_display=message.text.strip())
    await state.set_state(PeriodReport.date_to)
    await message.answer("Введите дату конца периода (например, 31.08.2026):")


@router.message(PeriodReport.date_to)
async def period_date_to(message: Message, state: FSMContext):
    d = _parse_date(message.text)
    if not d:
        await message.answer("Неверный формат. Введите дату (например, 31.08.2026):")
        return

    data = await state.get_data()
    date_from = data["date_from"]
    date_to = d.strftime("%Y-%m-%d")
    await state.clear()

    async with async_session() as session:
        workers = (await session.execute(
            select(User).where(User.role == "worker")
        )).scalars().all()

    if not workers:
        await message.answer("Работников нет.")
        return

    lines = []
    total = 0.0
    for worker in workers:
        earned, count, penalties, coeff = await _calc_worker_pay_period(worker, date_from, date_to)
        if count == 0 and penalties == 0:
            continue
        lines.append(
            f"• {worker.full_name} | {rank_name(worker.rank)}\n"
            f"  Задач: {count} | Коэфф: x{coeff}\n"
            f"  Штрафы: -{penalties:.0f}₽ | Итого: {earned:.0f}₽"
        )
        total += earned

    if not lines:
        await message.answer(f"За период {data['date_from_display']} — {message.text.strip()} данных нет.")
        return

    report = f"📊 Отчёт за период {data['date_from_display']} — {message.text.strip()}:\n\n"
    report += "\n\n".join(lines)
    report += f"\n\n{'='*30}\nОбщий итог: {total:.0f}₽"
    await message.answer(report)


# --- Статистика работника ---

@router.message(Command("my_stats"))
async def cmd_my_stats(message: Message):
    async with async_session() as session:
        user = (await session.execute(
            select(User).where(User.tg_id == message.from_user.id)
        )).scalar_one_or_none()

        if not user:
            await message.answer("Сначала зарегистрируйтесь: /start")
            return

    earned, confirmed_count, penalties, coeff = await _calc_worker_pay_monthly(user)

    await message.answer(
        f"📊 Ваша статистика (текущий месяц):\n\n"
        f"Ранг: {rank_name(user.rank)}\n"
        f"Подтверждённых задач: {confirmed_count}\n"
        f"Коэффициент: x{coeff}\n"
        f"Штрафы: -{penalties:.0f}₽\n"
        f"Заработано: {earned:.0f}₽"
    )


# --- Мои задачи ---

@router.message(Command("my_tasks"))
async def cmd_my_tasks(message: Message):
    async with async_session() as session:
        user = (await session.execute(
            select(User).where(User.tg_id == message.from_user.id)
        )).scalar_one_or_none()

        if not user:
            await message.answer("Сначала зарегистрируйтесь: /start")
            return

        apps = (await session.execute(
            select(TaskApplication)
            .options(selectinload(TaskApplication.task))
            .where(TaskApplication.user_id == user.id)
            .order_by(TaskApplication.id.desc())
        )).scalars().all()

    if not apps:
        await message.answer("У вас пока нет задач.")
        return

    status_labels = {
        ApplicationStatus.APPLIED.value: "📝 Записан",
        ApplicationStatus.PHOTO_START.value: "📸 Фото начала",
        ApplicationStatus.PHOTO_END.value: "📸 Фото конца",
        ApplicationStatus.CONFIRMED.value: "✅ Подтверждено",
        ApplicationStatus.NOT_SHOWED.value: "❌ Не пришёл",
        ApplicationStatus.PENALTY.value: "⚠️ Штраф",
    }

    lines = []
    for app in apps:
        status = status_labels.get(app.status, app.status)
        rate = app.task.rate_for_rank(user.rank)
        line = f"#{app.task.id} {app.task.title} — {app.task.date}\n   {status}"
        if app.status == ApplicationStatus.CONFIRMED.value:
            line += f" | {rate:.0f}₽"
        elif app.status == ApplicationStatus.PENALTY.value:
            line += f" | -{app.penalty_amount:.0f}₽"
        lines.append(line)

    await message.answer("📋 Мои задачи:\n\n" + "\n\n".join(lines))


# --- Профиль ---

@router.message(Command("profile"))
async def cmd_profile(message: Message):
    async with async_session() as session:
        user = (await session.execute(
            select(User).where(User.tg_id == message.from_user.id)
        )).scalar_one_or_none()

        if not user:
            await message.answer("Сначала зарегистрируйтесь: /start")
            return

    earned, confirmed_count, penalties, coeff = await _calc_worker_pay_monthly(user)
    reg_date = user.registered_at.strftime("%d.%m.%Y") if user.registered_at else "—"
    roles = ", ".join(user.job_roles) if user.job_roles else user.job_role

    await message.answer(
        f"👤 Мой профиль\n\n"
        f"Имя: {user.full_name}\n"
        f"Телефон: {user.phone or '—'}\n"
        f"Роли: {roles}\n"
        f"Ранг: {rank_name(user.rank)}\n"
        f"Дата регистрации: {reg_date}\n\n"
        f"📊 Статистика (текущий месяц):\n"
        f"Задач выполнено: {confirmed_count}\n"
        f"Коэффициент: x{coeff}\n"
        f"Штрафы: -{penalties:.0f}₽\n"
        f"Заработано: {earned:.0f}₽"
    )
